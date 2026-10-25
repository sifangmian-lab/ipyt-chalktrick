# -*- coding: utf-8 -*-
"""生成 2026-10-25 增补 DOCX（方案增补 + 模型报告更新/勘误）。

使用 DSH 自带 Python 的 python-docx；生成后可由 Word/COM 转 PDF。
运行：
  python code/tools/build_update_docs.py
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _set_font(run, name_cn="微软雅黑", size=10.5, bold=None):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name_cn)


def new_doc(title, subtitle):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    p = doc.add_heading(level=0)
    run = p.add_run(title)
    _set_font(run, size=16, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        run2 = p2.add_run(subtitle)
        _set_font(run2, size=10, bold=False)
        run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def h(doc, text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    _set_font(run, size=13 if level == 1 else 11.5, bold=True)


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_font(run, bold=bold)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        _set_font(run)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        _set_font(run, size=9.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            _set_font(run, size=9.5)


def build_plan_supplement(path):
    doc = new_doc("虚线戏法实验与理论建模方案（2026-10-25 增补）",
                  "v2 模型：弹性杆 + 手端正交弹簧；S4–S7 结果与 E3/E5/E6/E6A/E7/E8 增补规程。"
                  "与 Markdown 有效版冲突时，以 Markdown 与模型 v2 规格为准。")
    h(doc, "1. 模型 v2 更新摘要")
    bullets(doc, [
        "v1（归档）：固定铰轴 + 刚体棒 + 尖端法向罚弹簧 + tanh 摩擦；只算 N 下降段，无真正黏着；S1–S3 为 v1 数值基准。",
        "v2（目标）：3-DOF 集总模型 q=(x_H, y_H, θ)：手端 k_x, k_y（可加阻尼）、等效弯曲 k_θ=3EI/L、单边法向接触 k_n、切向 LuGre/spring-slider。",
        "痕端分类：黏着截断（v_rel→0，N 有限）/ 动力脱离（N→0）/ sprag-slip 冲击；不能默认“痕迹结束 = 压力归零”。",
        "关键文档：docs/研究记录/模型v2规格_2026-10-25.md、模型更新影响评审_2026-10-25.md、E6A截断判据执行清单_2026-10-25.md。",
    ])
    h(doc, "2. S4–S7/S5b 仿真结果摘要（2026-10-25 A2 预载修正后）")
    table(doc,
          ["编号", "内容", "主要结论"],
          [["S4", "v2 基线（spring-slider 本构）",
            "基线 A：N=3.969–4.000 N 平台（0.25% 涨落），I_c/I_plateau=1.002，有限强度截断；守恒 ∫I dξ=I_2nd。"],
           ["S5", "收敛性/正则化",
            "σ0、rtol、阻尼稳健；同一 k_n 双分辨率一致；罚接触无唯一 k_n→∞ 极限。"],
           ["S5b", "接触极限（刚性接触/LCP）",
            "t_stop=1.550 ms、I_c/I_p=1.000、约束残差 3e-13、双容差零差异；罚/刚性定量差 7%–15%。"],
           ["S6", "参数平面",
            "罚模型 k_n=2e4 下 (k_x,k_y) 全平面“平台+截断”；旧“k_y 主控”结论作废；低速划痕过短。"],
           ["S7", "本构对比",
            "LuGre-springslider 干净截断（1.002）；viscous LuGre 下降尾（0.646）；tanh（0.932）不用于截断预言。"]])
    h(doc, "3. 实验增补规程")
    table(doc,
          ["实验", "增补内容", "文档"],
          [["E3", "正交柔度模块 + k_x,k_y,k_θ/EI、阻尼；装置/人手分开标定",
            "docs/实验指导/E3_正交柔度模块与标定规程_2026-10-25.md"],
           ["E5", "痕端事件同步测量：N、切向力、v_rel、离面、声学；三层测量链",
            "docs/实验指导/E5_痕端事件测量链_2026-10-25.md"],
           ["E6", "逐划痕 I(ξ) 轮廓与痕端指标 I_plateau/I_c/R_c/S_c/ξ_c",
            "docs/实验指导/E6_痕端轮廓提取规程_2026-10-25.md"],
           ["E6A", "截断真实性判据；黏着/脱离/sprag-slip 判别",
            "docs/实验指导/E6A截断判据执行清单_2026-10-25.md"],
           ["E7/E8", "v2 反演链与普遍性；平台+截断指标",
            "docs/研究记录/E7_E8_v2更新_2026-10-25.md"]])
    h(doc, "4. 优先级与数据状态")
    bullets(doc, [
        "P0：用 E3 实测 k_x,k_y,k_θ 与 E1/E2 实测 μ_s,μ_k 替换占位参数；在此之前 v2 只作定性/趋势预言。",
        "P1：E5/E6/E6A 截断判据落地；S5b 刚性接触已给出 k_n 极限参考；S6b 角度/摩擦平面。",
        "P2：完整周期/极限环、E7 盲反演、E8 材料矩阵、DOCX/PDF 重生成。",
        "当前无 E1/E4 原始数据，也无截断原始图像/力数据；所有 v2 定量结果均为占位参数下的数值预言。",
        "原方案 DOCX/PDF 为历史版本；本文件为 2026-10-25 增补版。",
    ])
    doc.save(path)
    print("saved", path)


def build_report_supplement(path):
    doc = new_doc("粉笔 N 下降段模型：2026-10-25 更新与勘误",
                  "适用于《粉笔书写单周期压力下降段的建模与数值模拟》（2026-10-20）与 v2 模型更新。")
    h(doc, "1. I(ξ) 后处理勘误")
    para(doc, "原报告正文的 I(ξ) 定义是正确的（I·dξ=q·dt，∫I dξ=I_2nd），但数值实现用了 "
              "q·|dξ| 权重，量纲为功率而非 J/m；它多乘了一个 |v_rel|，导致图 2d 的“明显下降段”被放大、"
              "而原本要强调的折叠峰被压制。修正后名义工况（N0=4 N）：前向平台约 1.40→1.04 J/m，"
              "回划折叠峰约 2.07 J/m，回划尾到 0；∫I dξ/I_2nd=1.0000。")
    para(doc, "S2 的 μN·dt 权重在 v_rel→0 处发散，旧图中约 500 J/m 的“折叠尖峰”为伪影；"
              "S3 的分离点 F_f 掩膜已由 g>0 改为 Nn>1e-6·N0。")
    h(doc, "2. v1 边界与撤回")
    bullets(doc, [
        "v1（固定铰轴 + 刚体棒 + 法向罚弹簧 + tanh）无真正黏着；分离事件定义为 N=0，结构上不能产生有限强度截断。",
        "S1–S3 降级为 v1 数值基准；t*、I_2nd、ℓ 只在 v1 内有效，不能作为实验定量预言。",
        "原报告的图 2/预言③须配读勘误与 2026-10-25 重生成的 figures/结果图/划痕相对强度_时空依赖.png。",
    ])
    h(doc, "3. v2 模型与 S4–S7 结果")
    para(doc, "v2：弹性杆 + 手端正交弹簧 k_x,k_y,k_θ，单边法向接触，切向 LuGre/spring-slider。")
    table(doc,
          ["项目", "结果/状态"],
          [["S4 基线 A", "N=3.969–4.000 N（0.25% 涨落）；I_c/I_plateau=1.002（有限强度截断）；∫I/I2=0.99994"],
           ["S5", "σ0/rtol/阻尼稳健；同一 k_n 双分辨率一致；罚接触无唯一 k_n→∞ 极限"],
           ["S5b", "刚性接触（LCP）：t_stop=1.550 ms、I_c/I_p=1.000、约束残差 3e-13；罚/刚性差 7%–15%"],
           ["S6", "罚模型 (k_x,k_y) 全平面平台+截断；旧“k_y 主控”结论作废；低板速过短"],
           ["S7", "springslider 干净截断（1.002）；viscous LuGre 下降尾（0.646）；tanh（0.932）不用于截断预言"]])
    h(doc, "4. 结论")
    bullets(doc, [
        "“平台 + 有限强度截断”在 v2 的 spring-slider 基线中已复现，对应 stop point（v_rel→0、N 有限）。",
        "接触刚度 k_n、摩擦 σ1/σ2、k_x/k_y/k_θ 仍需 E3/E1/E2/E4 标定，当前数值只作定性/趋势。",
        "完整报告 PDF 未重建；本文件与 docs/模型报告/粉笔_N下降段_报告勘误_2026-10-25.md 共同构成更新说明。",
    ])
    doc.save(path)
    print("saved", path)


if __name__ == "__main__":
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    build_plan_supplement(os.path.join(
        root, "docs", "实验指导", "虚线戏法实验方案_2026-10-25增补.docx"))
    build_report_supplement(os.path.join(
        root, "docs", "模型报告", "粉笔_N下降段_模型更新与勘误_2026-10-25.docx"))
